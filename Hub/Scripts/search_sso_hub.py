from __future__ import annotations

import threading
import tkinter as tk
from pathlib import Path
from tkinter import messagebox, ttk
from openpyxl import load_workbook

ROOT = Path(r"G:\Dev\darkuniverse-training")
LINE_FOLDER = ROOT / "Shared" / "Input" / "Line Schedule"
SF_FOLDER = ROOT / "Shared" / "Input" / "SF Report"
PQS_OUTPUT = ROOT / "Reports" / "PQS Dark Universe Report" / "Output"
PQS_OPERATIONS = ROOT / "Reports" / "PQS Dark Universe Report" / "Input" / "Maximo" / "Operations"
PM_OUTPUT = ROOT / "Reports" / "PM Qualification Activity Tracker" / "Output"
DUE_SOON_DAYS = 15

COLORS={"navy":"#17365D","blue":"#2F75B5","light":"#EAF2F8","green":"#C6E0B4","yellow":"#FFE699","red":"#F4CCCC","gray":"#E7E6E6","white":"#FFFFFF","text":"#1F1F1F"}

def clean(v):
    if v is None:return ""
    t=str(v).strip()
    return t[:-2] if t.endswith(".0") and t[:-2].isdigit() else t

def norm(v):return " ".join(clean(v).casefold().split())
def newest(folder):
    files=[p for p in Path(folder).glob("*.xls*") if p.is_file() and not p.name.startswith("~$")]
    return max(files,key=lambda p:p.stat().st_mtime) if files else None

def find_header(ws, required, limit=40):
    best=(0,{})
    for r,row in enumerate(ws.iter_rows(min_row=1,max_row=min(ws.max_row,limit),values_only=True),1):
        mapping={norm(v):i for i,v in enumerate(row) if clean(v)}
        score=sum(any(norm(alias) in mapping for alias in aliases) for aliases in required.values())
        if score>best[0]:best=(score,(r,mapping))
    return best[1] if best[0] else (0,{})

def pick(mapping,*names):
    for name in names:
        if norm(name) in mapping:return mapping[norm(name)]
    return None

def value(row,mapping,*names):
    i=pick(mapping,*names)
    return clean(row[i]) if i is not None and i<len(row) else ""

def line_lookup(sso):
    path=newest(LINE_FOLDER); matches=[]
    if not path:return False,matches,"Unavailable"
    day_aliases={"sun":"Sun","sunday":"Sun","mon":"Mon","monday":"Mon","tue":"Tue","tues":"Tue","tuesday":"Tue","wed":"Wed","wednesday":"Wed","thur":"Thur","thurs":"Thur","thu":"Thur","thursday":"Thur","fri":"Fri","friday":"Fri","sat":"Sat","saturday":"Sat"}
    wb=load_workbook(path,read_only=True,data_only=True)
    try:
        for ws in wb.worksheets:
            rows=list(ws.iter_rows(values_only=True))
            day_columns={}; field_columns={}
            for header_index,row in enumerate(rows[:12]):
                for column_index,raw in enumerate(row):
                    header=norm(raw)
                    if header in day_aliases:day_columns[column_index]=day_aliases[header]
                    if header in {"technician","overstaff","overhire","id number","call sign"}:field_columns[column_index]=clean(raw)
            for row in rows:
                vals=[clean(v) for v in row]
                if sso not in vals:continue
                working=[];off=[]
                for column_index,day in day_columns.items():
                    marker=norm(vals[column_index]) if column_index<len(vals) else ""
                    if marker=="off":off.append(day)
                    else:working.append(day)
                details=[]
                for column_index,label in field_columns.items():
                    cell=vals[column_index] if column_index<len(vals) else ""
                    if not cell:continue
                    if norm(label) in {"supervisor","manager"}:continue
                    details.append(f"{label}: {cell}")
                if not any(sso in detail for detail in details):details.append(f"SSO: {sso}")
                matches.append({
                    "shift":ws.title,
                    "days":", ".join(working) if working else "No working days detected",
                    "details":" | ".join(details),
                })
    finally:wb.close()
    return bool(matches),matches,path.name

SF_ALIASES={"sso":["Employee ID","SSO","Labor ID"],"first":["First Name"],"last":["Last Name"],"display":["Display Name"],"item":["Item Title"],"days":["Days Until Due"],"org":["Org Description","Organization Description","Venue"],"park":["Park Description"],"dept":["Sub Division Description","Department"]}
def sf_lookup(sso):
    path=newest(SF_FOLDER)
    if not path:return None,[],"Unavailable"
    wb=load_workbook(path,read_only=True,data_only=True); person=None; items=[]; seen=set()
    try:
        for ws in wb.worksheets:
            hr,m=find_header(ws,SF_ALIASES)
            if not hr or pick(m,*SF_ALIASES["sso"]) is None:continue
            for row in ws.iter_rows(min_row=hr+1,values_only=True):
                if value(row,m,*SF_ALIASES["sso"])!=sso:continue
                if person is None:person={"first":"","last":"","display":"","org":set(),"park":set(),"dept":set()}
                for key in ("first","last","display"):
                    v=value(row,m,*SF_ALIASES[key]); person[key]=v or person[key]
                for key in ("org","park","dept"):
                    v=value(row,m,*SF_ALIASES[key]);
                    if v:person[key].add(v)
                d=value(row,m,*SF_ALIASES["days"])
                try:days=float(d)
                except:continue
                if days>DUE_SOON_DAYS:continue
                days=int(days) if days.is_integer() else days
                item=value(row,m,*SF_ALIASES["item"]); org=value(row,m,*SF_ALIASES["org"])
                status="Overdue" if days<0 else "Due Soon"; key=(item,days,status,org)
                if key not in seen:seen.add(key);items.append((item,days,status,org))
            if person:break
    finally:wb.close()
    items.sort(key=lambda x:(0 if x[2]=="Overdue" else 1,x[1],x[0].casefold()))
    return person,items,path.name

def pm_lookup(sso):
    results=[]; sources=[]
    required={"sso":["SSO"],"qual":["Qualification"],"cycle":["Cycle Number"],"start":["Cycle Start Date"],"due":["Cycle Due Date","Annual Due Date"],"pm":["PM Completion Date","Valid PM Date"],"days":["Days Remaining"],"status":["Cycle Result","Status"],"task":["Task Number"],"source":["Task Source","Evidence Source"]}
    for path in PM_OUTPUT.glob("*.xlsx"):
        if path.name.startswith("~$"):continue
        wb=load_workbook(path,read_only=True,data_only=True)
        try:
            if "PM Status" not in wb.sheetnames:continue
            ws=wb["PM Status"];hr,m=find_header(ws,required)
            for row in ws.iter_rows(min_row=hr+1,values_only=True):
                if value(row,m,"SSO")!=sso:continue
                record=(value(row,m,"Qualification"),value(row,m,"Cycle Number"),value(row,m,"Cycle Start Date"),value(row,m,"Cycle Due Date","Annual Due Date"),value(row,m,"PM Completion Date","Valid PM Date"),value(row,m,"Days Remaining"),value(row,m,"Cycle Result","Status"),value(row,m,"Task Number"),value(row,m,"Task Source","Evidence Source"))
                results.append(record);sources.append(path.name)
        finally:wb.close()
    return results,sorted(set(sources))

def pqs_operations_lookup(sso):
    rows=[]
    path=newest(PQS_OPERATIONS)
    if not path:return rows,""
    wb=load_workbook(path,read_only=True,data_only=True)
    try:
        for ws in wb.worksheets:
            ws.reset_dimensions();header_row=0;header_map={}
            for row_number,row in enumerate(ws.iter_rows(min_row=1,max_row=20,values_only=True),1):
                mapping={norm(v):i for i,v in enumerate(row) if clean(v)}
                if {"trainee","wo #","module","status"}.issubset(mapping):
                    header_row=row_number;header_map=mapping;break
            if not header_row:continue
            trainee_col=header_map["trainee"];location_col=header_map.get("location");work_order_col=header_map["wo #"]
            module_col=header_map["module"];status_col=header_map["status"];total_col=header_map.get("total tasks");completed_col=header_map.get("completed tasks")
            current_sso="";current_location="";current_work_order=""
            for row in ws.iter_rows(min_row=header_row+1,values_only=True):
                trainee=clean(row[trainee_col]) if trainee_col<len(row) else ""
                if trainee:
                    current_sso=trainee
                    current_location=clean(row[location_col]) if location_col is not None and location_col<len(row) else ""
                    current_work_order=clean(row[work_order_col]) if work_order_col<len(row) else ""
                if current_sso!=sso:continue
                module=clean(row[module_col]) if module_col<len(row) else ""
                status=clean(row[status_col]).upper() if status_col<len(row) else ""
                if not module or status not in {"INPRG","COMP"}:continue
                total=clean(row[total_col]) if total_col is not None and total_col<len(row) else ""
                completed=clean(row[completed_col]) if completed_col is not None and completed_col<len(row) else ""
                progress=f"{completed}/{total}" if completed or total else ""
                rows.append((current_location,"",module,progress,status,current_work_order,path.name))
    finally:wb.close()
    return rows,path.name if rows else ""


def pqs_lookup(sso):
    rows=[];sources=[]
    path=newest(PQS_OUTPUT)
    if path:
        wb=load_workbook(path,read_only=True,data_only=True)
        try:
            for ws in wb.worksheets:
                required={"sso":["SSO"],"name":["Name","Technician"],"position":["Position"],"venue":["Venue"],"shift":["Shift"]}
                hr,m=find_header(ws,required)
                if not hr or pick(m,"SSO") is None:continue
                headers=next(ws.iter_rows(min_row=hr,max_row=hr,values_only=True))
                for row in ws.iter_rows(min_row=hr+1,values_only=True):
                    if value(row,m,"SSO")!=sso:continue
                    base={"sheet":ws.title,"position":value(row,m,"Position"),"venue":value(row,m,"Venue"),"shift":value(row,m,"Shift")}
                    found=False
                    for i,h in enumerate(headers):
                        title=clean(h);upper=title.upper()
                        if not title or not ("MOD" in upper or "STATUS" in upper or "PROGRESS" in upper or "REMAIN" in upper or "REVIEW" in upper):continue
                        v=clean(row[i]) if i<len(row) else ""
                        if v:rows.append((base["venue"],base["shift"],title,v,"","",ws.title));found=True
                    if not found:rows.append((base["venue"],base["shift"],"Summary",base["position"],"","",ws.title))
            if rows:sources.append(path.name)
        finally:wb.close()
    operation_rows,operation_source=pqs_operations_lookup(sso)
    rows.extend(operation_rows)
    if operation_source:sources.append(operation_source)
    return rows,sorted(set(sources))

class App(tk.Toplevel):
    def __init__(self, parent=None):
        super().__init__(parent);self.title("Search SSO - Dark Universe Reporting Hub");self.geometry("1050x700");self.minsize(900,620);self.configure(bg="#F4F6F8")
        self.sso=tk.StringVar();self.name=tk.StringVar(value="Search a Team Member");self.sub=tk.StringVar(value="Enter an SSO / Employee ID to view current training information.");self.sources=tk.StringVar(value="")
        self.build();self.bind("<Return>",lambda e:self.start());self.entry.focus_set()
    def build(self):
        hero=tk.Frame(self,bg=COLORS["navy"],padx=24,pady=18);hero.pack(fill="x")
        tk.Label(hero,text="SEARCH SSO",font=("Segoe UI",22,"bold"),bg=COLORS["navy"],fg="white").pack(anchor="w")
        tk.Label(hero,text="Dark Universe training, qualifications, PM cycles, and SuccessFactors",font=("Segoe UI",10),bg=COLORS["navy"],fg="#D9EAF7").pack(anchor="w")
        body=tk.Frame(self,bg="#F4F6F8",padx=20,pady=16);body.pack(fill="both",expand=True)
        search=tk.Frame(body,bg="white",highlightthickness=1,highlightbackground="#CDD5DF",padx=14,pady=14);search.pack(fill="x")
        search.columnconfigure(0,weight=1);self.entry=ttk.Entry(search,textvariable=self.sso,font=("Segoe UI",12));self.entry.grid(row=0,column=0,sticky="ew",padx=(0,10));self.btn=ttk.Button(search,text="Search",command=self.start);self.btn.grid(row=0,column=1)
        tk.Label(body,textvariable=self.name,font=("Segoe UI",16,"bold"),bg="#F4F6F8",fg=COLORS["text"]).pack(anchor="w",pady=(14,2))
        tk.Label(body,textvariable=self.sub,font=("Segoe UI",10),bg="#F4F6F8",fg="#505A66").pack(anchor="w")
        self.banner=tk.Label(body,text="Ready",anchor="w",font=("Segoe UI",10,"bold"),padx=12,pady=8,bg=COLORS["gray"],fg="#404040");self.banner.pack(fill="x",pady=(10,10))
        self.tabs=ttk.Notebook(body);self.tabs.pack(fill="both",expand=True)
        self.trees={}
        self.addtab("line","Profile & Line Schedule",("Shift / Sheet","Days Worked","Line Schedule Details"),(170,240,760))
        self.addtab("pqs","PQS Modules",("Venue","Shift","Module / Field","Result / Progress","Maximo Status","Work Order","Source"),(140,90,260,180,110,140,190))
        self.addtab("pm","PM Cycles",("Qualification","Cycle","Cycle Start","Due Date","PM Date","Days Remaining","Cycle Status","Task","Evidence"),(110,65,105,105,105,105,170,110,150))
        self.addtab("sf","SuccessFactors",("Item Title","Days Until Due","Status","Organization / Venue"),(520,120,110,300))
        tk.Label(body,textvariable=self.sources,font=("Segoe UI",8),bg="#F4F6F8",fg="#606A75",anchor="w").pack(fill="x",pady=(8,0))
    def addtab(self,key,title,cols,widths):
        f=ttk.Frame(self.tabs,padding=8);self.tabs.add(f,text=title);f.rowconfigure(0,weight=1);f.columnconfigure(0,weight=1)
        t=ttk.Treeview(f,columns=cols,show="headings",selectmode="browse")
        for c,w in zip(cols,widths):t.heading(c,text=c);t.column(c,width=w,minwidth=60,anchor="w")
        t.grid(row=0,column=0,sticky="nsew");ys=ttk.Scrollbar(f,orient="vertical",command=t.yview);ys.grid(row=0,column=1,sticky="ns");xs=ttk.Scrollbar(f,orient="horizontal",command=t.xview);xs.grid(row=1,column=0,sticky="ew");t.configure(yscrollcommand=ys.set,xscrollcommand=xs.set)
        for tag,color in [("Good",COLORS["green"]),("Good / Completed",COLORS["green"]),("Due Soon",COLORS["yellow"]),("Overdue",COLORS["red"]),("Requalification Needed",COLORS["red"]),("Task Data Missing","#E4DFEC")]:t.tag_configure(tag,background=color)
        self.trees[key]=t
    def start(self):
        s=self.sso.get().strip()
        if not s:messagebox.showwarning("Search SSO","Enter an SSO / Employee ID.");return
        self.btn.configure(state="disabled");self.banner.configure(text="Searching current data...",bg=COLORS["gray"]);threading.Thread(target=self.worker,args=(s,),daemon=True).start()
    def worker(self,s):
        try:
            sf,sfitems,sfsrc=sf_lookup(s);du,line,linesrc=line_lookup(s);pqs,pqssrc=pqs_lookup(s) if du else ([],[]);pm,pmsrc=pm_lookup(s) if du else ([],[])
            self.after(0,self.show,s,sf,sfitems,sfsrc,du,line,linesrc,pqs,pqssrc,pm,pmsrc)
        except Exception as e:self.after(0,self.fail,str(e))
    def clear(self):
        for t in self.trees.values():
            for i in t.get_children():t.delete(i)
    def show(self,s,sf,sfitems,sfsrc,du,line,linesrc,pqs,pqssrc,pm,pmsrc):
        self.clear();first=sf.get("first","") if sf else "";last=sf.get("last","") if sf else "";display=sf.get("display","") if sf else ""
        self.name.set(f"{first or display or 'Unknown'} {last}".strip());org=", ".join(sorted(sf["org"])) if sf else "";sf_summary = "All Good" if sf and not sfitems else (f"{len(sfitems)} outstanding item(s)" if sf else "SF not found")
        self.sub.set(f"SSO: {s}"+(f"  |  {org}" if org else "")+f"  |  SuccessFactors: {sf_summary}")
        if du:self.banner.configure(text="Dark Universe Team Member - Full Dark Universe and SuccessFactors information shown.",bg=COLORS["green"],fg="#1F4E24")
        elif sf:self.banner.configure(text="Not currently listed on the Dark Universe Line Schedule - SuccessFactors information only.",bg=COLORS["yellow"],fg="#7F6000")
        else:self.banner.configure(text="SSO not found in the current Line Schedule or SuccessFactors workbook.",bg=COLORS["gray"],fg="#404040")
        for r in line:self.trees["line"].insert("","end",values=(r["shift"],r["days"],r["details"]))
        for r in pqs:self.trees["pqs"].insert("","end",values=r,tags=(r[3],))
        for r in pm:self.trees["pm"].insert("","end",values=r,tags=(r[6],))
        if sf and not sfitems:
            self.trees["sf"].insert("","end",values=("All Good","", "Good", "No outstanding items with Days Until Due <= 15"),tags=("Good",))
        else:
            for r in sfitems:self.trees["sf"].insert("","end",values=r,tags=(r[2],))
        self.tabs.tab(1,text=f"PQS Modules ({len(pqs)})");self.tabs.tab(2,text=f"PM Cycles ({len(pm)})");self.tabs.tab(3,text=f"SuccessFactors ({len(sfitems) if sfitems else (1 if sf else 0)})")
        self.sources.set(f"Sources - Line Schedule: {linesrc} | SF: {sfsrc} | PQS: {', '.join(pqssrc) or '-'} | PM: {', '.join(pmsrc) or '-'}")
        self.btn.configure(state="normal");self.entry.focus_set()
    def fail(self,e):self.btn.configure(state="normal");self.banner.configure(text="Search failed.",bg=COLORS["red"]);messagebox.showerror("Search SSO",e)

def launch_search_sso(parent=None):
    window = App(parent)
    if parent is not None:
        window.transient(parent)
    window.focus_force()
    return window

if __name__=="__main__":
    root=tk.Tk();root.withdraw();window=launch_search_sso(root);root.mainloop()
